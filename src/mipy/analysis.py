"""One call from a data table to results, figures and a report.

>>> import mipy
>>> res = mipy.analyze("my_data.csv")          # or a DataFrame
>>> res.show()                                  # figures and tables (inline in Jupyter)
>>> res.completed                               # the completed table

``analyze`` checks the table, hides a share of the observed values under
several missing-data mechanisms, imputes them with each method, scores the
imputations against the hidden values, imputes the real missing values
with the method that did best, and writes tables (CSV and Excel), figures
(PNG and PDF), the completed data and an HTML report into one folder.
A stage that fails is reported in the warnings and the rest still runs.
"""
from __future__ import annotations

import datetime as _dt
import html
import re
import sys
import time
import traceback
import warnings
from functools import partial
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Union

import numpy as np
import pandas as pd

from .context import ImputeContext
from .data import DataError, PreparedData, prepare_data
from .imputers import METHODS, available

PRESETS = {
    "quick": ["Mean", "LOCF", "KNN", "softImpute", "MinProb", "QRILC", "MICE-Long"],
    "standard": ["Mean", "Median", "LOCF", "KNN", "softImpute", "MICE-Bayes", "sklearn-II",
                 "MinProb", "QRILC", "MICE-Long", "MICE-LongMNAR"],
    "complete": [k for k, v in METHODS.items() if v.main],
    "all": list(METHODS),
}
MECH_ORDER = ("MCAR", "MAR", "MNAR", "LOD")
SCORE_VERSION = "analyze-score-v1"
_PIP = {"miceforest": "pip install miceforest lightgbm", "hyperimpute": "pip install hyperimpute"}


# ---------------------------------------------------------------- scoring (runs in workers)
def _score_task(rec, X_true, mask, patient):
    """Reduce an imputation run to its scores; keeps no matrices."""
    from .evaluate import evaluate_per_protein, structure_ratio
    if not rec.get("ok"):
        return {"ok": False, "error": rec.get("error", ""), "seconds": rec.get("seconds", 0.0)}
    out = rec["imputed"]
    p = X_true.shape[1]
    truth = np.where(mask, X_true, np.nan)
    pp = evaluate_per_protein(truth, out, mask, names=list(range(p)))
    vec = np.full(p, np.nan)
    if len(pp):
        vec[pp["protein"].to_numpy(dtype=int)] = pp["nrmse"].to_numpy()
    res = {"ok": True, "error": "", "seconds": rec["seconds"], "pp": vec,
           "nRMSE": float(pp["nrmse"].median()) if len(pp) else np.nan,
           "r": float(pp["r"].median()) if len(pp) else np.nan,
           "bias": float(pp["bias"].median()) if len(pp) else np.nan,
           "pct_below_1": float(100 * (pp["nrmse"] < 1).mean()) if len(pp) else np.nan,
           "sd_ratio": (structure_ratio(X_true, out, mask, patient) if patient is not None else np.nan),
           "coverage": np.nan, "width": np.nan, "m": rec.get("m", 1)}
    draws = rec.get("draws")
    if draws and len(draws) > 1:
        from .mi import predictive_coverage
        cov = predictive_coverage(truth, draws, mask)
        res["coverage"], res["width"] = 100 * cov["coverage"], cov["width"]
    return res


# ---------------------------------------------------------------- helpers
def _in_notebook() -> bool:
    """True in Jupyter, JupyterLab, VS Code notebooks and Colab (not in
    Spyder's console, which cannot show HTML tables)."""
    if "spyder_kernels" in sys.modules:
        return False
    try:
        from IPython import get_ipython
        ip = get_ipython()
        return ip is not None and hasattr(ip, "kernel")
    except Exception:
        return False


def _safe(name) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(name))


def _resolve_methods(methods, prep: PreparedData):
    if methods is None:
        methods = "standard"
    if isinstance(methods, str):
        key = methods.lower()
        if key in PRESETS:
            wanted = PRESETS[key]
        else:
            wanted = [s.strip() for s in methods.split(",") if s.strip()]
    else:
        wanted = list(methods)
    lookup = {k.lower(): k for k in METHODS}
    chosen, skipped = [], {}
    for w in wanted:
        name = lookup.get(str(w).lower())
        if name is None:
            raise ValueError(f"unknown method {w!r}. Choose from: {', '.join(METHODS)}; or a preset: "
                             f"{', '.join(PRESETS)}")
        info = METHODS[name]
        if name in chosen:
            continue
        if not available(name):
            skipped[name] = f"needs the package '{info.needs}' ({_PIP.get(info.needs, 'pip install ' + str(info.needs))})"
        elif info.longitudinal and not prep.longitudinal:
            skipped[name] = "needs repeated visits per patient (an ID column with 2+ rows per patient)"
        else:
            chosen.append(name)
    return chosen, skipped


def _plausible_mechanism(X, mechanisms):
    """Spearman correlation between a feature's median observed value and
    its share of missing values; strongly negative suggests that low values
    are lost (MNAR, as below a detection limit)."""
    from scipy.stats import spearmanr
    M = np.isnan(X)
    pf = M.mean(0)
    med = np.nanmedian(np.where(M, np.nan, X), axis=0) if (~M).any() else np.full(X.shape[1], np.nan)
    g = np.isfinite(med)
    rho = np.nan
    if g.sum() >= 5 and np.ptp(pf[g]) > 0:
        rho = float(spearmanr(med[g], pf[g])[0])
    if "MNAR" in mechanisms and np.isfinite(rho) and rho <= -0.3:
        return "MNAR", rho, (f"features with lower values have more missing values (Spearman rho = "
                             f"{rho:.2f}), as when values fall below a detection limit")
    pick = "MAR" if "MAR" in mechanisms else mechanisms[0]
    rtxt = f"Spearman rho = {rho:.2f}" if np.isfinite(rho) else "rho not defined"
    return pick, rho, (f"missingness is not strongly tied to how high a feature's values are "
                       f"({rtxt}), so {pick} is the working assumption")


# ---------------------------------------------------------------- result object
class Result:
    """What :func:`analyze` returns. Main attributes:

    summary          accuracy of every method under every mechanism (DataFrame)
    overview         mean nRMSE, methods x mechanisms (DataFrame)
    recommendation   dict: best method per mechanism, and the one used for export
    decision         the same as a table
    completed        the input table with the missing values imputed (DataFrame)
    completed_draws  the m completed tables of the multiple-imputation method (list)
    survival         exploratory Cox screen (DataFrame) or None
    tables, figures  every table (DataFrame) and figure (PNG path) by name
    warnings         anything that did not go as planned
    report_path      the HTML report; out_dir the output folder
    """

    def __init__(self):
        self.prepared: Optional[PreparedData] = None
        self.data_summary = pd.DataFrame()
        self.scores = pd.DataFrame()
        self.summary = pd.DataFrame()
        self.overview = pd.DataFrame()
        self.per_feature = pd.DataFrame()
        self.masking = pd.DataFrame()
        self.decision = pd.DataFrame()
        self.recommendation: dict = {}
        self.completed: Optional[pd.DataFrame] = None
        self.completed_method: Optional[str] = None
        self.completed_draws: List[pd.DataFrame] = []
        self.mi_method: Optional[str] = None
        self.fmi = pd.DataFrame()
        self.survival: Optional[pd.DataFrame] = None
        self.tables: dict = {}
        self.figures: dict = {}
        self.figure_notes: dict = {}
        self.files: List[Path] = []
        self.methods: List[str] = []
        self.skipped: dict = {}
        self.warnings: List[str] = []
        self.notes: List[str] = []
        self.settings: dict = {}
        self.versions: dict = {}
        self.out_dir: Optional[Path] = None
        self.report_path: Optional[Path] = None
        self.seconds = 0.0
        self.log_lines: List[str] = []
        self._details: List[str] = []
        self._completed_analysis = None          # analysis scale, analysis row order
        self._draws_analysis = None

    # ------------------------------------------------------------ text
    def text_summary(self) -> str:
        lines = ["MIPY analysis"]
        if self.prepared is not None:
            p = self.prepared
            lines.append(f"  data: {len(p.X)} rows, {p.n_patients} patients, {len(p.features)} features, "
                         f"{100 * np.isnan(p.X).mean():.2f}% missing")
        for mech, v in self.recommendation.get("best", {}).items():
            lines.append(f"  best under {mech:4s}: {v['method']} (nRMSE {v['nRMSE']:.3f})")
        if self.completed_method:
            lines.append(f"  completed data: {self.completed_method}"
                         + (f"; {len(self.completed_draws)} completions from {self.mi_method}"
                            if self.mi_method else ""))
        if self.survival is not None:
            lines.append(f"  survival screen: {int((self.survival['q'] < 0.10).sum())} features with FDR q < 0.10")
        if self.warnings:
            lines.append(f"  warnings: {len(self.warnings)} (see res.warnings or the report)")
        if self.report_path:
            lines.append(f"  report: {self.report_path}")
        return "\n".join(lines)

    def __repr__(self):
        return self.text_summary()

    def _repr_html_(self):
        from .report import result_html_summary
        return result_html_summary(self)

    # ------------------------------------------------------------ display
    def show(self, figures: Optional[Sequence[str]] = None):
        """Show the main tables and every figure: inline in Jupyter, or as
        text plus file paths in a terminal."""
        names = list(figures) if figures else list(self.figures)
        if _in_notebook():
            from IPython.display import HTML, Image, display
            display(HTML(self._repr_html_()))
            for n in names:
                if n in self.figures:
                    display(HTML(f"<h4 style='margin:14px 0 2px'>{html.escape(n)}</h4>"
                                 f"<p style='margin:0 0 6px;color:#444'>{html.escape(self.figure_notes.get(n, ''))}</p>"))
                    from .report import png_width
                    display(Image(filename=str(self.figures[n]), width=min(900, png_width(self.figures[n]))))
        else:
            print(self.text_summary())
            if len(self.overview):
                print("\nmean nRMSE (lower is better; 1 = feature mean)")
                print(self.overview.round(3).to_string())
            print("\nfigures:")
            for n in names:
                if n in self.figures:
                    print(f"  {n}: {self.figures[n]}")
            if self.report_path:
                print(f"\nopen the report in a browser: {self.report_path}")

    def open_report(self):
        """Open the HTML report in the default web browser."""
        import webbrowser
        if self.report_path:
            webbrowser.open(Path(self.report_path).resolve().as_uri())


# ---------------------------------------------------------------- stages
class _Stage:
    """Runs a block; on an error, records a warning and lets the analysis go on."""

    def __init__(self, res: Result, name: str, log):
        self.res, self.name, self.log = res, name, log

    def __enter__(self):
        return self

    def __exit__(self, et, e, tb):
        if e is None or isinstance(e, (KeyboardInterrupt, SystemExit)):
            return False
        msg = f"{self.name} could not be completed ({type(e).__name__}: {e})"
        self.res.warnings.append(msg)
        self.res._details.append(f"{self.name}\n" + "".join(traceback.format_exception(et, e, tb)))
        self.log("  ! " + msg)
        return True


def _versions():
    out = {"mipy": None, "python": sys.version.split()[0]}
    from . import __version__
    out["mipy"] = __version__
    for mod in ("numpy", "pandas", "scipy", "sklearn", "matplotlib", "miceforest", "lightgbm",
                "hyperimpute"):
        try:
            m = __import__(mod)
            out[mod] = getattr(m, "__version__", "?")
        except Exception:
            pass
    return out


# ---------------------------------------------------------------- main
def analyze(data, id_col=None, time_col=None, features=None, clinical=None, survival=None,
            out: Union[str, Path] = "mipy_results", methods: Union[str, Sequence[str]] = "standard",
            mechanisms: Iterable[str] = ("MCAR", "MAR", "MNAR"), repeats: int = 3,
            rate: float = 0.2, strength: float = 2.0, m: int = 5, log2: Union[str, bool] = "auto",
            export: Union[str, None] = "auto", n_workers: Union[int, str] = 1, seed: int = 2026,
            max_missing: float = 0.9, sheet=0, layout: str = "wide", feature_col=None,
            value_col=None, resume: bool = True, show: bool = True, verbose: bool = True,
            settings: Optional[dict] = None) -> Result:
    """Benchmark imputation methods on your table, impute it, and report.

    Parameters
    ----------
    data : file path (CSV, Excel, TSV, Parquet) or DataFrame; one row per
        patient visit, like the GSE65622 files (ID, visit, features, and
        optional clinical and survival columns).
    id_col, time_col, features, clinical, survival : see
        :func:`mipy.data.prepare_data`; found automatically when not given.
    out : folder for the results (created if needed).
    methods : "quick", "standard" (default), "complete", "all", or a list
        of names from ``mipy.list_methods()``.
    mechanisms : masking designs, from "MCAR", "MAR", "MNAR", "LOD".
    repeats : maskings per mechanism (more = smaller Monte Carlo error).
    rate : share of the observed values hidden in each masking.
    strength : slope of the MAR and MNAR masking (larger = stronger).
    m : completions per multiple-imputation method.
    log2 : "auto", True or False (see prepare_data).
    export : "auto" (best method under the most plausible mechanism), a
        method name, or None for no completed data.
    n_workers : processes for the imputations; 1 (default, safest) or
        "auto" (CPU count - 1, at most 4).
    resume : reuse finished runs saved in ``out/.mipy_cache``.
    show : at the end, show the results (inline in Jupyter).
    settings : other method settings, e.g. {"mice_iter": 10, "gamma": 1.5}.
    """
    t_start = time.time()
    res = Result()
    T0 = _dt.datetime.now()

    def log(msg):
        line = str(msg)
        res.log_lines.append(line)
        if verbose:
            try:
                print(line, flush=True)
            except UnicodeEncodeError:                  # consoles that cannot show e.g. Greek letters
                print(line.encode("ascii", "replace").decode("ascii"), flush=True)

    # ------------------------------------------------ arguments
    mechs = []
    for mm in ([mechanisms] if isinstance(mechanisms, str) else list(mechanisms)):
        for part in str(mm).split(","):
            u = part.strip().upper()
            if not u:
                continue
            if u not in MECH_ORDER:
                raise ValueError(f"unknown mechanism {part!r}; choose from {', '.join(MECH_ORDER)}")
            if u not in mechs:
                mechs.append(u)
    if not mechs:
        raise ValueError("give at least one mechanism, e.g. mechanisms=('MCAR', 'MAR', 'MNAR')")
    mechs = [u for u in MECH_ORDER if u in mechs]
    repeats = int(repeats)
    if repeats < 1:
        raise ValueError("repeats must be 1 or more")
    if not 0.01 <= float(rate) <= 0.6:
        raise ValueError("rate must be between 0.01 and 0.6 (share of observed values hidden)")
    m = int(m)
    if m < 2:
        raise ValueError("m must be 2 or more (completions for multiple imputation)")
    settings = dict(settings or {})
    known = set(ImputeContext.__dataclass_fields__) - {"seed", "patient", "time", "schema", "m", "extra"}
    bad = set(settings) - known
    if bad:
        raise ValueError(f"unknown setting(s) {sorted(bad)}; allowed: {sorted(known)}")
    if n_workers == "auto":
        from .runner import default_workers
        n_workers = default_workers()
    n_workers = max(1, int(n_workers))

    lookup = {k.lower(): k for k in METHODS}
    if not (isinstance(methods, str) and methods.lower() in PRESETS) and methods is not None:
        for w in ([s for s in methods.split(",")] if isinstance(methods, str) else methods):
            if str(w).strip().lower() not in lookup:
                raise ValueError(f"unknown method {w!r}. Choose from: {', '.join(METHODS)}; or a "
                                 f"preset: {', '.join(PRESETS)}")
    if export and str(export).lower() != "auto" and str(export).lower() not in lookup:
        raise ValueError(f"export: unknown method {export!r}")

    out = Path(out).expanduser()
    res.out_dir = out
    res.versions = _versions()
    log(f"MIPY {res.versions['mipy']}: analysis started {T0:%Y-%m-%d %H:%M}; results go to {out.resolve()}")

    # ------------------------------------------------ 1. data
    log("1/7 checking the data")
    prep = prepare_data(data, id_col=id_col, time_col=time_col, features=features, clinical=clinical,
                        survival=survival, log2=log2, max_missing=max_missing, sheet=sheet,
                        layout=layout, feature_col=feature_col, value_col=value_col)
    res.prepared = prep
    res.data_summary = prep.summary()
    res.notes += prep.notes
    res.warnings += prep.warnings
    for n in prep.notes:
        log("    " + n)
    for w in prep.warnings:
        log("  ! " + w)
    X = prep.X
    n, p = X.shape

    methods_run, skipped = _resolve_methods(methods, prep)
    res.methods, res.skipped = methods_run, skipped
    for k, v in skipped.items():
        log(f"  - {k} skipped: {v}")
    if not methods_run:
        raise ValueError("none of the chosen methods can run on these data; see the messages above")
    figdir, tabdir, compdir = out / "figures", out / "tables", out / "completed_data"
    for d in (out, figdir, tabdir, compdir):              # created only once the input is known to be usable
        d.mkdir(parents=True, exist_ok=True)
    data_name = Path(data).name if isinstance(data, (str, Path)) else "DataFrame"
    res.settings = {"data": data_name, "methods": ", ".join(methods_run), "mechanisms": ", ".join(mechs),
                    "repeats": repeats, "share hidden (rate)": rate, "masking slope (strength)": strength,
                    "completions (m)": m, "seed": seed, "workers": n_workers,
                    "analysis scale": "log2(x + 1)" if prep.log2 else "as given",
                    "export": export if export else "none", **settings}

    # ------------------------------------------------ 2. masking
    log(f"2/7 hiding {int(round(100 * rate))}% of the observed values: {', '.join(mechs)} x {repeats} repeat(s)")
    from .masking import Amputer
    from .runner import CheckpointStore, fingerprint, run_one, run_tasks
    from .schema import build_variable_schema
    min_obs = int(min(12, max(3, n // 5)))
    amper = Amputer(X, time=prep.time, clinical=prep.clinical, rate=rate, strength=strength,
                    min_obs=min_obs)
    amps, mrows, zvals = {}, [], {}
    for rep in range(1, repeats + 1):
        for mech in mechs:
            j = MECH_ORDER.index(mech)
            amp = amper.ampute(mech, seed + 1000 * (rep - 1) + 101 * (j + 1))
            amps[(rep, mech)] = amp
            mrows.append({"repeat": rep, "mechanism": mech, **amper.summary(amp)})
            if rep == 1:
                zvals[mech] = amper.Z[amp["mask"]]
    res.masking = pd.DataFrame(mrows)

    # ------------------------------------------------ 3. imputation runs
    pat_sd = prep.patient if prep.longitudinal else None
    ctx_kw = dict(seed=seed, patient=prep.patient, time=prep.time, m=m, **settings)
    base_ctx = ImputeContext(**ctx_kw)
    data_fp = fingerprint(X, "data", (repr(None if prep.patient is None else prep.patient.tolist()),
                                      repr(None if prep.time is None else prep.time.tolist())))
    tasks = []
    for (rep, mech), amp in amps.items():
        sch = build_variable_schema(amp["X"], base_ctx.max_predictors, base_ctx.min_cor)
        ctx = base_ctx.with_(schema=sch)
        for name in methods_run:
            info = METHODS[name]
            sett = (info.rev, seed, m, ctx.mice_iter, ctx.knn_k, ctx.max_predictors, ctx.min_cor,
                    ctx.gain_epochs, ctx.hyper_inner_iter, ctx.gamma, ctx.mi_holdout, data_fp,
                    SCORE_VERSION, repr(sorted(settings.items())))
            tasks.append({"key": (rep, mech, name), "name": name, "X": amp["X"], "ctx": ctx,
                          "ckpt": "t_" + fingerprint(amp["X"], name, sett),
                          "post": partial(_score_task, X_true=X, mask=amp["mask"], patient=pat_sd)})
    store = CheckpointStore(out / ".mipy_cache", fresh=not resume)
    total = len(tasks)
    log(f"3/7 imputing: {len(methods_run)} methods x {len(amps)} maskings = {total} runs"
        f"{f' on {n_workers} workers' if n_workers > 1 else ''}")
    done = {"k": 0}

    def on_done(t, rec, loaded):
        done["k"] += 1
        rep, mech, name = t["key"]
        st = ("saved earlier" if loaded else f"{rec.get('seconds', 0):.1f} s") if rec.get("ok") \
            else f"FAILED: {rec.get('error', '')[:120]}"
        log(f"    [{done['k']:>{len(str(total))}}/{total}] {mech:4s} repeat {rep}  {name:14s} {st}")

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        results = run_tasks(tasks, store, n_workers=n_workers, on_done=on_done, log=log)

    # ------------------------------------------------ 4. scores
    log("4/7 scoring")
    rows, pp_rows = [], []
    feats = prep.features
    for (rep, mech, name), r in results.items():
        rows.append({"repeat": rep, "mechanism": mech, "method": name, "ok": bool(r.get("ok")),
                     "nRMSE": r.get("nRMSE", np.nan), "r": r.get("r", np.nan),
                     "bias": r.get("bias", np.nan), "% features nRMSE < 1": r.get("pct_below_1", np.nan),
                     "SD ratio": r.get("sd_ratio", np.nan), "coverage (%)": r.get("coverage", np.nan),
                     "interval width": r.get("width", np.nan), "seconds": r.get("seconds", np.nan),
                     "error": r.get("error", "")})
        if r.get("ok") and r.get("pp") is not None:
            pp_rows.append(pd.DataFrame({"repeat": rep, "mechanism": mech, "method": name,
                                         "feature": feats, "nRMSE": r["pp"]}))
    scores = pd.DataFrame(rows)
    scores["mechanism"] = pd.Categorical(scores["mechanism"], [u for u in MECH_ORDER if u in mechs])
    scores = scores.sort_values(["mechanism", "repeat", "nRMSE"]).reset_index(drop=True)
    scores["mechanism"] = scores["mechanism"].astype(str)
    res.scores = scores
    failed = scores[~scores["ok"]]
    for (name, err), g in failed.groupby(["method", "error"]):
        res.warnings.append(f"{name} failed in {len(g)} of {int((scores['method'] == name).sum())} runs: {err[:200]}")
    ok = scores[scores["ok"]].dropna(subset=["nRMSE"])
    if not scores["ok"].any():
        res.warnings.append("no imputation run finished; see the errors above")
    elif ok.empty:
        res.warnings.append("the methods could not be scored: each feature needs at least 3 hidden "
                            "values, and this table is too small for that (use more rows, or a "
                            "larger rate, e.g. rate=0.3)")
    from .evaluate import mcse_mean
    best_count = (ok.loc[ok.groupby(["repeat", "mechanism"])["nRMSE"].idxmin()]
                  .groupby(["mechanism", "method"]).size() if len(ok) else pd.Series(dtype=int))
    srows = []
    for (mech, name), g in scores.groupby(["mechanism", "method"], sort=False):
        go = g[g["ok"]]
        srows.append({"mechanism": mech, "method": name, "family": METHODS[name].family,
                      "nRMSE": go["nRMSE"].mean() if len(go) else np.nan,
                      "MCSE": mcse_mean(go["nRMSE"]) if len(go) > 1 else np.nan,
                      "best in": f"{int(best_count.get((mech, name), 0))} of {repeats}",
                      "% features nRMSE < 1": go["% features nRMSE < 1"].mean() if len(go) else np.nan,
                      "r": go["r"].mean() if len(go) else np.nan,
                      "bias": go["bias"].mean() if len(go) else np.nan,
                      "SD ratio": go["SD ratio"].mean() if len(go) else np.nan,
                      "coverage (%)": go["coverage (%)"].mean() if len(go) and go["coverage (%)"].notna().any() else np.nan,
                      "interval width": go["interval width"].mean() if len(go) and go["interval width"].notna().any() else np.nan,
                      "seconds": go["seconds"].mean() if len(go) else np.nan,
                      "runs ok": f"{len(go)} of {len(g)}"})
    summary = pd.DataFrame(srows)
    if len(summary):
        summary["rank"] = summary.groupby("mechanism")["nRMSE"].rank(method="min")
        summary["_m"] = summary["mechanism"].map({u: i for i, u in enumerate(MECH_ORDER)})
        summary = summary.sort_values(["_m", "nRMSE"]).drop(columns="_m").reset_index(drop=True)
    res.summary = summary
    if len(summary):
        res.overview = (summary.pivot(index="method", columns="mechanism", values="nRMSE")
                        .reindex(columns=mechs))
        res.overview = res.overview.loc[res.overview.mean(axis=1).sort_values().index]
    if pp_rows:
        pp = pd.concat(pp_rows, ignore_index=True)
        res.per_feature = (pp.groupby(["mechanism", "method", "feature"], sort=False)["nRMSE"].mean()
                           .reset_index())

    # ------------------------------------------------ 5. recommendation and export
    log("5/7 choosing a method and imputing the real missing values")
    plaus, rho, why = _plausible_mechanism(X, mechs)
    best = {}
    drows = []
    for mech in mechs:
        s = summary[(summary["mechanism"] == mech) & summary["nRMSE"].notna()].sort_values("nRMSE")
        if s.empty:
            continue
        b = s.iloc[0]
        best[mech] = {"method": b["method"], "nRMSE": float(b["nRMSE"])}
        smi = s[s["method"].map(lambda k: METHODS[k].multiple)]
        mi_b = smi.iloc[0] if len(smi) else None
        drows.append({"mechanism": mech, "best method": b["method"], "nRMSE": b["nRMSE"],
                      "best in": b["best in"],
                      "runner-up": s.iloc[1]["method"] if len(s) > 1 else "",
                      "runner-up nRMSE": s.iloc[1]["nRMSE"] if len(s) > 1 else np.nan,
                      "best multiple-imputation method": mi_b["method"] if mi_b is not None else "",
                      "its coverage (%)": mi_b["coverage (%)"] if mi_b is not None else np.nan})
    res.decision = pd.DataFrame(drows)
    res.recommendation = {"best": best, "plausible mechanism": plaus, "rho": rho, "why": why}

    exp_name = None
    if export:
        if str(export).lower() == "auto":
            ranked = (summary[(summary["mechanism"] == plaus) & summary["nRMSE"].notna()]
                      .sort_values("nRMSE")["method"].tolist() if len(summary) else [])
            if not ranked and len(summary):
                ranked = summary.groupby("method")["nRMSE"].mean().dropna().sort_values().index.tolist()
            if not ranked:                                  # nothing could be scored: robust defaults
                ranked = [k for k in ("softImpute", "KNN", "Mean") if k in methods_run] or ["softImpute", "Mean"]
                res.recommendation.update({"plausible mechanism": "none",
                                           "why": "the methods could not be scored on this table, so a "
                                                  "robust default was used"})
        else:
            lookup = {k.lower(): k for k in METHODS}
            exp_name = lookup.get(str(export).lower())
            if exp_name is None:
                raise ValueError(f"export: unknown method {export!r}")
            ranked = [exp_name]
            res.recommendation.update({"export chosen by user": True,
                                       "export reason": f"chosen with export='{exp_name}'"})
        with _Stage(res, "Imputing the real missing values", log):
            sch = build_variable_schema(X, base_ctx.max_predictors, base_ctx.min_cor)
            ctx = base_ctx.with_(schema=sch)
            ran = {}
            for name in ranked:
                if METHODS[name].longitudinal and not prep.longitudinal:
                    continue
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    rec = run_one(name, X, ctx)
                ran[name] = rec
                if rec["ok"]:
                    exp_name = name
                    break
                res.warnings.append(f"{name} could not impute the full table ({rec['error'][:150]}); "
                                    f"trying the next method")
            if exp_name and ran.get(exp_name, {}).get("ok"):
                rec = ran[exp_name]
                res.completed_method = exp_name
                if "export reason" not in res.recommendation:
                    first = next(iter(ran))
                    res.recommendation["export reason"] = (
                        res.recommendation["why"] if res.recommendation.get("plausible mechanism") == "none"
                        else f"the best method under {plaus}, because {res.recommendation['why']}"
                        + ("" if first == exp_name else f"; {first} was better but could not complete "
                                                        f"the full table"))
                res.completed = prep.completed_table(rec["imputed"])
                fn = compdir / f"completed_{_safe(exp_name)}.csv"
                res.completed.to_csv(fn, index=False)
                res.files.append(fn)
                log(f"    completed table ({exp_name}): {fn.name}")
                res.recommendation["export"] = exp_name
                # multiple imputation
                if METHODS[exp_name].multiple and rec.get("draws"):
                    mi_name, draws = exp_name, rec["draws"]
                else:
                    mi_rank = [k for k in ranked if METHODS[k].multiple]
                    if str(export).lower() != "auto":
                        s = summary[(summary["mechanism"] == plaus) & summary["nRMSE"].notna()]
                        mi_rank = [k for k in s.sort_values("nRMSE")["method"] if METHODS[k].multiple]
                    mi_name, draws = None, None
                    for k in mi_rank:
                        if METHODS[k].longitudinal and not prep.longitudinal:
                            continue
                        with warnings.catch_warnings():
                            warnings.simplefilter("ignore")
                            r2 = run_one(k, X, ctx)
                        if r2["ok"] and r2.get("draws"):
                            mi_name, draws = k, r2["draws"]
                            break
                if mi_name:
                    res.mi_method = mi_name
                    res.recommendation["export_mi"] = mi_name
                    for l, D in enumerate(draws, 1):
                        tab = prep.completed_table(D)
                        res.completed_draws.append(tab)
                        fn = compdir / f"completed_{_safe(mi_name)}_draw{l}.csv"
                        tab.to_csv(fn, index=False)
                        res.files.append(fn)
                    log(f"    {len(draws)} completions for multiple imputation ({mi_name})")
                    if np.isnan(X).any():
                        from .mi import fmi_summary
                        f = fmi_summary(draws, X)
                        res.fmi = pd.DataFrame([{"method": mi_name, "m": f["m"],
                                                 "FMI, mean of imputed cells": f["FMI"],
                                                 "median FMI, feature means": f["median FMI, protein means"],
                                                 "IQR FMI, feature means": f["IQR FMI, protein means"]}])
                    res._draws_analysis = draws
                res._completed_analysis = rec["imputed"]
            elif not ran:
                res.warnings.append("no method was available to impute the full table")

    # ------------------------------------------------ 6. survival screen
    if prep.survival:
        log("6/7 exploratory survival screen")
        with _Stage(res, "Survival screen", log):
            comps = res._draws_analysis or (
                [res._completed_analysis] if res._completed_analysis is not None else None)
            if comps is None:
                res.notes.append("survival screen skipped: no completed data")
            else:
                res.survival = _survival_screen(prep, comps, res)
                log("    " + res.notes[-1])
    else:
        log("6/7 survival screen: no survival columns, skipped")

    # ------------------------------------------------ 7. figures, tables, report
    log("7/7 writing figures, tables and the report")
    _figures(res, prep, zvals, rate, figdir, log)
    _tables(res, tabdir, log)
    res.seconds = time.time() - t_start
    with _Stage(res, "Report", log):
        from .report import write_report
        res.report_path = write_report(res, out / "report.html")
        res.files.append(res.report_path)
    _readme(res, out)
    (out / "run_log.txt").write_text("\n".join(res.log_lines + ["", "warnings:"] + res.warnings
                                               + ["", "details:"] + res._details), encoding="utf-8")
    log(f"done in {res.seconds / 60:.1f} min. Open {res.report_path or out}")
    if show:
        try:
            res.show()
        except Exception as e:                          # display problems never stop the analysis
            log(f"  (could not display the results here: {e}; open the report instead)")
    return res


# ---------------------------------------------------------------- survival
def _survival_screen(prep: PreparedData, comps, res: Result):
    from .survival import pooled_univariate_cox
    tcol, ecol = prep.survival
    if prep.patient is not None:
        first = np.unique(prep.patient, return_index=True)[1]
    else:
        first = np.arange(len(prep.X))
    rows = prep.order[first]
    T = pd.to_numeric(prep.table[tcol].iloc[rows], errors="coerce").to_numpy(dtype=float)
    E = pd.to_numeric(prep.table[ecol].iloc[rows], errors="coerce").to_numpy(dtype=float)
    ok = np.isfinite(T) & np.isfinite(E) & (T > 0)
    n_ev = int(E[ok].sum())
    if ok.sum() < 20 or n_ev < 10:
        res.notes.append(f"survival screen skipped: {int(ok.sum())} patients with a usable time and "
                         f"{n_ev} events (needs at least 20 patients and 10 events)")
        return None
    C = [np.asarray(c, dtype=float)[first][ok] for c in comps]
    tab = pooled_univariate_cox(T[ok], E[ok].astype(int), C, prep.features)
    tab = tab.rename(columns={"protein": "feature"})
    tab.insert(2, "HR per SD", np.exp(tab["log_hr"]))
    tab = tab.sort_values("p").reset_index(drop=True)
    res.notes.append(f"survival screen: baseline sample of {int(ok.sum())} patients, {n_ev} events, "
                     f"univariate Cox per feature"
                     + (f", pooled over {len(C)} completions with Rubin's rules" if len(C) > 1 else ""))
    return tab


# ---------------------------------------------------------------- figures and tables
FIG_NOTES = {
    "missingness": "Where values are missing. A: share missing per feature. B: by visit. C: the "
                   "pattern (black = missing). D: if low-level features miss more often (negative "
                   "rho), values are probably lost below a detection limit (MNAR).",
    "masking": "The benchmark hides observed values whose true value is known. Under MCAR the hidden "
               "values are typical (z around 0); under MNAR and LOD they are low, as for values "
               "below a detection limit.",
    "accuracy": "Main result. nRMSE of each method on the hidden values (median across features, "
                "mean over repeats; bars = 95% Monte Carlo interval). 1.0 is what filling in each "
                "feature's mean achieves; lower is better.",
    "nrmse_per_feature": "Spread of the error across features (each feature's nRMSE averaged over "
                         "repeats). A method with a low median but a long box fails on some features.",
    "sd_ratio": "Within-patient SD of the completed data divided by that of the true values. Near 1: "
                "visit-to-visit changes are kept. Below 1: imputed values are too flat (e.g. mean "
                "imputation), which hides change over time.",
    "coverage": "For multiple-imputation methods: share of hidden values inside the 95% interval "
                "built from the m completions. Near 95% means the uncertainty is honest; far below "
                "means the completions are too similar.",
    "runtime": "Average seconds per imputation of the whole table.",
    "imputed_vs_observed": "The real missing values as imputed by the chosen method, against the "
                           "observed values. A: distributions. B: per feature, how far the imputed "
                           "values sit from the observed mean (negative = imputed lower, as expected "
                           "when values are missing because they are low).",
    "survival": "Exploratory: univariate Cox model per feature on the baseline sample, pooled over "
                "the completions. Red points pass a 10% false discovery rate. Not a validated "
                "prognostic model.",
}


def _figures(res: Result, prep: PreparedData, zvals, rate, figdir, log):
    try:
        import matplotlib  # noqa: F401
    except ImportError:
        res.warnings.append("figures were not drawn: matplotlib is not installed (pip install matplotlib)")
        return
    from . import plots
    colors = plots.method_colors(res.methods)
    jobs = [("missingness", lambda: plots.missingness(prep.X, prep.features, prep.time if prep.time_col else None)),
            ("masking", lambda: plots.masking(zvals, rate))]
    if len(res.summary) and res.summary["nRMSE"].notna().any():
        jobs.append(("accuracy", lambda: plots.accuracy(res.summary, colors)))
    if len(res.per_feature):
        jobs.append(("nrmse_per_feature", lambda: plots.protein_nrmse(res.per_feature, colors)))
    if len(res.summary) and res.summary["SD ratio"].notna().any():
        jobs.append(("sd_ratio", lambda: plots.sd_ratio(res.summary)))
    if len(res.summary) and res.summary["coverage (%)"].notna().any():
        jobs.append(("coverage", lambda: plots.coverage(res.summary)))
    if len(res.scores) and res.scores["ok"].any():
        jobs.append(("runtime", lambda: plots.runtime(res.scores, colors)))
    if res._completed_analysis is not None and np.isnan(prep.X).any():
        jobs.append(("imputed_vs_observed",
                     lambda: plots.imputed_vs_observed(prep.X, res._completed_analysis, res.completed_method)))
    if res.survival is not None and len(res.survival):
        jobs.append(("survival", lambda: plots.survival_volcano(res.survival)))
    for k, (name, make) in enumerate(jobs, 1):
        fname = f"fig{k:02d}_{name}"
        with _Stage(res, f"Figure '{name}'", log):
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                fig = make()
                res.figures[name] = plots.save_figure(fig, figdir, fname)
            res.figure_notes[name] = FIG_NOTES.get(name, "")
            res.files += [res.figures[name], res.figures[name].with_suffix(".pdf")]


def _round(df):
    out = df.copy()
    for c in out.columns:
        if pd.api.types.is_float_dtype(out[c]):
            out[c] = out[c].round(4)
    return out


def _tables(res: Result, tabdir, log):
    tabs = {"data_summary": res.data_summary, "accuracy_summary": res.summary,
            "accuracy_overview": res.overview.reset_index() if len(res.overview) else pd.DataFrame(),
            "decision": res.decision, "all_runs": res.scores, "masking": res.masking,
            "per_feature_nrmse": res.per_feature, "fmi": res.fmi,
            "methods": _methods_table(res)}
    if res.survival is not None:
        tabs["survival_screen"] = res.survival
    res.tables = {k: v for k, v in tabs.items() if isinstance(v, pd.DataFrame) and len(v)}
    with _Stage(res, "Writing tables", log):
        for k, v in res.tables.items():
            fn = tabdir / f"{k}.csv"
            _round(v).to_csv(fn, index=False)
            res.files.append(fn)
    try:
        import openpyxl  # noqa: F401
        with _Stage(res, "Writing the Excel workbook", log):
            fn = tabdir / "mipy_results.xlsx"
            with pd.ExcelWriter(fn) as xw:
                for k, v in res.tables.items():
                    _round(v).to_excel(xw, sheet_name=k[:31], index=False)
            res.files.append(fn)
    except ImportError:
        res.notes.append("Excel workbook not written (pip install openpyxl); the CSV tables are complete")


def _methods_table(res: Result):
    rows = []
    for name in res.methods + list(res.skipped):
        i = METHODS[name]
        rows.append({"method": name, "family": i.family, "package / function": f"{i.package}: {i.function}",
                     "multiple imputation": "yes" if i.multiple else "no",
                     "uses patients": "yes" if i.longitudinal else "no",
                     "peer-reviewed description": "yes" if i.peer_reviewed else "no",
                     "status": "run" if name in res.methods else f"skipped: {res.skipped[name]}"})
    return pd.DataFrame(rows)


def _readme(res: Result, out: Path):
    lines = ["MIPY results", "============", "",
             "report.html                    open in a web browser: all results with explanations",
             "figures/                       every figure as PNG (for slides) and PDF (for papers)",
             "tables/                        every table as CSV, and all of them in mipy_results.xlsx",
             "  accuracy_summary.csv         nRMSE, MCSE, SD ratio, coverage per method and mechanism",
             "  decision.csv                 best method under each mechanism",
             "  per_feature_nrmse.csv        error of each method on each feature",
             "completed_data/               your table with the missing values filled in",
             "run_log.txt                    what was done, warnings, error details",
             ".mipy_cache/                   finished runs; delete it to start from scratch", ""]
    if res.completed_method:
        lines.append(f"Completed data: {res.completed_method}"
                     + (f"; multiple imputation: {res.mi_method} ({len(res.completed_draws)} files "
                        f"*_draw1..{len(res.completed_draws)}.csv). Analyse each completed file and "
                        f"pool the estimates with Rubin's rules (mipy.mi.pool_estimates)." if res.mi_method else ""))
    (out / "README.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
